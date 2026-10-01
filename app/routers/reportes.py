import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import APIRouter, Depends, HTTPException, status, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import datetime, timezone
import pytz
from typing import List, Dict, Any, Optional
from .. import models, schemas, crud
from ..database import get_db
from ..timezone import (
    LIMA_TZ,
    obtener_ahora_lima,
    convertir_a_lima,
    formatear_fecha_lima,
    convertir_lima_a_utc_naive
)
from .auth import get_current_admin_user
from .tickets import estructurar_info_creador

router = APIRouter(
    prefix="/api/reportes",
    tags=["Reportes y Dashboards Gerenciales"]
)

dashboard_router = APIRouter(
    prefix="/api/dashboard",
    tags=["Dashboard Gerencial"]
)

# Tiempos SLA en horas según criticidad (HU18)
SLA_LIMITES_HORAS = {
    "Alta": 24,
    "Alto": 24,
    "Crítico": 24,
    "Critico": 24,
    "Media": 72,
    "Medio": 72,
    "Baja": 168,
    "Bajo": 168
}


def calcular_metrica_ticket(ticket: models.Ticket, ahora: Optional[datetime] = None) -> Dict[str, Any]:
    """
    Calcula el estado y tiempo restante del SLA para un ticket individual (HU18).
    Garantiza que tanto la fecha de creación del ticket como la variable de la
    'hora actual' se encuentren estrictamente en la misma zona horaria (America/Lima UTC-5)
    para eliminar el error matemático de las 'horas extra'.
    
    Aplica la nueva regla matemática de tercios:
    - Verde (success): Si ha transcurrido menos de un tercio (< 33.3%) del tiempo total.
    - Amarillo (warning): Si ha transcurrido entre un tercio y dos tercios (>= 33.3% y < 66.6%) del tiempo total.
    - Rojo (danger): Si ha transcurrido más de dos tercios (>= 66.6%) del tiempo total, o si ya está vencido.
    """
    if ahora is None:
        ahora = obtener_ahora_lima()
    elif ahora.tzinfo is None:
        ahora = pytz.utc.localize(ahora).astimezone(LIMA_TZ)
    else:
        ahora = ahora.astimezone(LIMA_TZ)

    criticidad = ticket.criticidad or "Media"
    sla_limite = SLA_LIMITES_HORAS.get(criticidad, 72)
    
    fecha_creacion = convertir_a_lima(ticket.fecha_creacion) or ahora

    tiempo_transcurrido = max(0.0, (ahora - fecha_creacion).total_seconds() / 3600.0)
    horas_restantes = max(0.0, sla_limite - tiempo_transcurrido)
    pct_transcurrido = (tiempo_transcurrido / sla_limite) * 100.0 if sla_limite > 0 else 100.0
    
    if ticket.estado == "Resuelto":
        # Considerado cumplido
        estado_sla = "Cumplido"
        esta_vencido = False
        color_sla = "success"
    else:
        if tiempo_transcurrido >= sla_limite or horas_restantes <= 0:
            estado_sla = "Vencido"
            esta_vencido = True
            color_sla = "danger"
        elif pct_transcurrido >= (200.0 / 3.0):  # >= 66.6% del tiempo total
            estado_sla = "En Riesgo"
            esta_vencido = False
            color_sla = "danger"
        elif pct_transcurrido >= (100.0 / 3.0):  # >= 33.3% y < 66.6% del tiempo total
            estado_sla = "En Atención"
            esta_vencido = False
            color_sla = "warning"
        else:  # < 33.3% del tiempo total
            estado_sla = "Normal"
            esta_vencido = False
            color_sla = "success"

    es_autoatencion = False
    if ticket.comentario_tecnico:
        com_lower = ticket.comentario_tecnico.lower()
        if "video" in com_lower or "autoatencion" in com_lower or "microaprendizaje" in com_lower:
            es_autoatencion = True

    return {
        "sla_limite_horas": sla_limite,
        "tiempo_transcurrido_horas": round(tiempo_transcurrido, 1),
        "horas_restantes": round(horas_restantes, 1),
        "porcentaje_transcurrido": round(pct_transcurrido, 1),
        "estado_sla": estado_sla,
        "color_sla": color_sla,
        "esta_vencido": esta_vencido,
        "es_autoatencion": es_autoatencion
    }


def generar_resumen_ejecutivo(db: Session) -> Dict[str, Any]:
    """Genera métricas consolidadas para la Jefatura de TI en hora local Lima."""
    tickets = db.query(models.Ticket).all()
    ahora = obtener_ahora_lima()

    total_tickets = len(tickets)
    abiertos = 0
    en_progreso = 0
    resueltos = 0
    
    por_criticidad = {"Alta": 0, "Media": 0, "Baja": 0}
    por_sede: Dict[str, int] = {}
    vencidos_sla = 0
    en_riesgo_sla = 0
    dentro_sla = 0
    criticos_alta = 0
    autoatencion_count = 0

    lista_tickets_reporte = []

    for t in tickets:
        # Estado
        if t.estado == "Abierto":
            abiertos += 1
        elif t.estado == "En Progreso":
            en_progreso += 1
        elif t.estado == "Resuelto":
            resueltos += 1

        # Criticidad
        crit_normalizada = "Media"
        if t.criticidad in ["Alta", "Alto", "Crítico", "Critico"]:
            crit_normalizada = "Alta"
            por_criticidad["Alta"] += 1
            if t.estado != "Resuelto":
                criticos_alta += 1
        elif t.criticidad in ["Baja", "Bajo"]:
            crit_normalizada = "Baja"
            por_criticidad["Baja"] += 1
        else:
            por_criticidad["Media"] += 1

        # Sede
        sede_nombre = t.sede or "Sin Sede Asignada"
        por_sede[sede_nombre] = por_sede.get(sede_nombre, 0) + 1

        # SLA
        sla_info = calcular_metrica_ticket(t, ahora)
        if sla_info["esta_vencido"]:
            vencidos_sla += 1
        else:
            dentro_sla += 1

        if sla_info["estado_sla"] == "En Riesgo":
            en_riesgo_sla += 1

        if sla_info["es_autoatencion"]:
            autoatencion_count += 1

        creador_info = estructurar_info_creador(db, t)
        fc_lima = convertir_a_lima(t.fecha_creacion)
        lista_tickets_reporte.append({
            "id": t.id,
            "descripcion": t.descripcion,
            "estado": t.estado,
            "criticidad": crit_normalizada,
            "sede": t.sede or "-",
            "piso": t.piso or "-",
            "fecha_creacion": fc_lima.isoformat() if fc_lima else None,
            "solicitante": creador_info.nombre or t.correo_solicitante or "Usuario",
            "correo_solicitante": t.correo_solicitante or creador_info.correo or "-",
            "telefono": creador_info.telefono or "-",
            "anexo": creador_info.anexo or "-",
            "estado_sla": sla_info["estado_sla"],
            "color_sla": sla_info["color_sla"],
            "horas_restantes": sla_info["horas_restantes"],
            "sla_limite_horas": sla_info["sla_limite_horas"],
            "tiempo_transcurrido_horas": sla_info["tiempo_transcurrido_horas"],
            "porcentaje_transcurrido": sla_info["porcentaje_transcurrido"],
            "es_autoatencion": sla_info["es_autoatencion"]
        })


    # Cumplimiento SLA porcentaje
    porcentaje_sla = 100.0
    if total_tickets > 0:
        porcentaje_sla = round(((total_tickets - vencidos_sla) / total_tickets) * 100, 1)

    tasa_resolucion = 0.0
    if total_tickets > 0:
        tasa_resolucion = round((resueltos / total_tickets) * 100, 1)

    return {
        "kpis": {
            "total_tickets": total_tickets,
            "abiertos": abiertos,
            "en_progreso": en_progreso,
            "resueltos": resueltos,
            "criticos_pendientes": criticos_alta,
            "cumplimiento_sla_porcentaje": porcentaje_sla,
            "tasa_resolucion_porcentaje": tasa_resolucion,
            "tickets_vencidos_sla": vencidos_sla,
            "tickets_en_riesgo_sla": en_riesgo_sla,
            "tickets_dentro_sla": dentro_sla,
            "autoatencion_resueltos": autoatencion_count
        },
        "distribucion": {
            "por_criticidad": por_criticidad,
            "por_sede": por_sede,
            "por_estado": {
                "Abierto": abiertos,
                "En Progreso": en_progreso,
                "Resuelto": resueltos
            }
        },
        "tickets": sorted(lista_tickets_reporte, key=lambda x: x["id"], reverse=True)
    }


# =========================================================================
# ENDPOINTS PROTEGIDOS EXCLUSIVOS PARA JEFATURA DE TI (HU20 - 403 FORBIDDEN)
# =========================================================================

@router.get("/dashboard")
def get_reportes_dashboard(
    db: Session = Depends(get_db),
    admin_user: models.Usuario = Depends(get_current_admin_user)
):
    """
    HU20: Endpoint exclusivo para Jefatura de TI.
    Retorna métricas consolidadas, KPIs de SLA y distribución operativa.
    Protegido con validación estricta de rol 403 Forbidden.
    """
    return generar_resumen_ejecutivo(db)


@router.get("/kpis")
def get_reportes_kpis(
    db: Session = Depends(get_db),
    admin_user: models.Usuario = Depends(get_current_admin_user)
):
    """Retorna los indicadores clave de rendimiento (KPIs) ejecutivos."""
    resumen = generar_resumen_ejecutivo(db)
    return resumen["kpis"]


@router.get("/tickets")
def get_reportes_tickets(
    sede: Optional[str] = Query(None),
    criticidad: Optional[str] = Query(None),
    estado: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    admin_user: models.Usuario = Depends(get_current_admin_user)
):
    """Retorna la lista de tickets para auditoría ejecutiva con filtros opcionales."""
    resumen = generar_resumen_ejecutivo(db)
    tickets = resumen["tickets"]

    if sede:
        tickets = [t for t in tickets if (t.get("sede") or "").lower() == sede.lower()]
    if criticidad:
        tickets = [t for t in tickets if (t.get("criticidad") or "").lower() == criticidad.lower()]
    if estado:
        tickets = [t for t in tickets if (t.get("estado") or "").lower() == estado.lower()]

    return tickets


@router.get("/exportar/excel")
def exportar_reportes_excel(
    fecha_inicio: Optional[str] = Query(None, description="Fecha de inicio (YYYY-MM-DD)"),
    fecha_fin: Optional[str] = Query(None, description="Fecha de fin (YYYY-MM-DD)"),
    db: Session = Depends(get_db),
    admin_user: models.Usuario = Depends(get_current_admin_user)
):
    """
    HU15.3: Exportación nativa estructurada a Excel con dos hojas:
    - Hoja 1 ("Resumen Ejecutivo"): Hoja activa con KPIs (Horas ahorradas, MTTR, FCR, Reabiertos),
      Top 5 Incidentes Recurrentes y Resumen Mensual del período.
    - Hoja 2 ("Data Detallada"): Tabla completa de todas las incidencias del período filtrado.
    """
    from .analytics import calcular_metricas_analytics

    # 1. Obtener métricas calculadas de Analytics para el período
    analytics_data = calcular_metricas_analytics(
        db=db,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        anio=None
    )
    kpis = analytics_data.get("kpis", {})
    top_5 = analytics_data.get("graficos", {}).get("top_incidentes_recurrentes", [])
    resumen_mensual = analytics_data.get("resumen_mensual", [])

    # 2. Consultar incidencias detalladas para la Hoja 2
    query = (
        db.query(models.Ticket, models.Usuario)
        .outerjoin(models.Usuario, models.Ticket.usuario_id == models.Usuario.id)
    )

    if fecha_inicio:
        try:
            dt_inicio_local = datetime.strptime(fecha_inicio.strip(), "%Y-%m-%d")
            dt_inicio_utc = convertir_lima_a_utc_naive(dt_inicio_local)
            query = query.filter(models.Ticket.fecha_creacion >= dt_inicio_utc)
        except ValueError:
            pass

    if fecha_fin:
        try:
            dt_fin_local = datetime.strptime(fecha_fin.strip(), "%Y-%m-%d").replace(hour=23, minute=59, second=59)
            dt_fin_utc = convertir_lima_a_utc_naive(dt_fin_local)
            query = query.filter(models.Ticket.fecha_creacion <= dt_fin_utc)
        except ValueError:
            pass

    results = query.order_by(models.Ticket.id.desc()).all()
    ahora = obtener_ahora_lima()

    wb = Workbook()

    # -------------------------------------------------------------------------
    # ESTILOS COMUNES OPENPYXL
    # -------------------------------------------------------------------------
    font_titulo = Font(name="Calibri", size=14, bold=True, color="1A1E23")
    font_subtitulo = Font(name="Calibri", size=10, italic=True, color="57606A")
    font_seccion_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_seccion_header = PatternFill(start_color="1A1E23", end_color="1A1E23", fill_type="solid")

    font_tabla_header = Font(name="Calibri", size=11, bold=True, color="1F2328")
    fill_tabla_header = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")

    font_negrita = Font(name="Calibri", size=11, bold=True)
    font_normal = Font(name="Calibri", size=11)

    border_fino = Border(
        left=Side(style='thin', color='D0D7DE'),
        right=Side(style='thin', color='D0D7DE'),
        top=Side(style='thin', color='D0D7DE'),
        bottom=Side(style='thin', color='D0D7DE')
    )

    # -------------------------------------------------------------------------
    # HOJA 1: RESUMEN EJECUTIVO (Hoja activa)
    # -------------------------------------------------------------------------
    ws_resumen = wb.active
    ws_resumen.title = "Resumen Ejecutivo"

    # Encabezado principal del informe
    ws_resumen.append(["HelpStream Analytics - Resumen Ejecutivo"])
    ws_resumen.cell(row=1, column=1).font = font_titulo

    periodo_str = f"Período: {fecha_inicio or 'Inicio'} al {fecha_fin or 'Actual'}" if (fecha_inicio or fecha_fin) else "Período: Consolidado Histórico Completo"
    ws_resumen.append([periodo_str, "", f"Generado: {ahora.strftime('%Y-%m-%d %H:%M:%S')} (Hora Lima UTC-5)"])
    ws_resumen.cell(row=2, column=1).font = font_subtitulo
    ws_resumen.cell(row=2, column=3).font = font_subtitulo
    ws_resumen.append([])  # Espacio

    # a) SECCIÓN: 4 KPIS PRINCIPALES
    row_kpi_sec = ws_resumen.max_row + 1
    ws_resumen.append(["INDICADORES CLAVE DE RENDIMIENTO (KPIS)", "", ""])
    c_sec1 = ws_resumen.cell(row=row_kpi_sec, column=1)
    c_sec1.font = font_seccion_header
    c_sec1.fill = fill_seccion_header

    ws_resumen.append(["Indicador / Métrica", "Valor", "Detalle Operativo"])
    row_kpi_head = ws_resumen.max_row
    for c in range(1, 4):
        cell = ws_resumen.cell(row=row_kpi_head, column=c)
        cell.font = font_tabla_header
        cell.fill = fill_tabla_header
        cell.border = border_fino

    filas_kpis = [
        ("Horas-Hombre Ahorradas (HU15.1)", f"{kpis.get('horas_ahorradas', 0.0)} hrs", f"Autoatención con microaprendizaje ({kpis.get('tickets_autoatencion', 0)} tickets resueltos)"),
        ("MTTR (Tiempo Medio de Resolución)", f"{kpis.get('mttr_horas', 0.0)} hrs", "Promedio de atención efectiva de incidencias"),
        ("Resolución en Primer Contacto (FCR)", f"{kpis.get('fcr_porcentaje', 0.0)}%", "Tasa de cierre directo sin escalamiento"),
        ("Ratio de Tickets Reabiertos", f"{kpis.get('ratio_reabiertos_porcentaje', 0.0)}%", f"{kpis.get('tickets_reabiertos', 0)} incidencias con reapertura registrada")
    ]

    for nom, val, det in filas_kpis:
        ws_resumen.append([nom, val, det])
        curr_row = ws_resumen.max_row
        ws_resumen.cell(row=curr_row, column=1).font = font_normal
        ws_resumen.cell(row=curr_row, column=1).border = border_fino
        c2 = ws_resumen.cell(row=curr_row, column=2)
        c2.font = font_negrita
        c2.alignment = Alignment(horizontal="center")
        c2.border = border_fino
        ws_resumen.cell(row=curr_row, column=3).font = font_normal
        ws_resumen.cell(row=curr_row, column=3).border = border_fino

    ws_resumen.append([])  # Espacio

    # b) SECCIÓN: TOP 5 INCIDENTES RECURRENTES
    row_top_sec = ws_resumen.max_row + 1
    ws_resumen.append(["TOP 5 INCIDENTES RECURRENTES", "", ""])
    c_sec2 = ws_resumen.cell(row=row_top_sec, column=1)
    c_sec2.font = font_seccion_header
    c_sec2.fill = fill_seccion_header

    ws_resumen.append(["Categoría / Título", "Cantidad de Casos", "% del Total"])
    row_top_head = ws_resumen.max_row
    for c in range(1, 4):
        cell = ws_resumen.cell(row=row_top_head, column=c)
        cell.font = font_tabla_header
        cell.fill = fill_tabla_header
        cell.border = border_fino

    total_t = kpis.get("total_tickets", len(results))
    if not top_5:
        ws_resumen.append(["Sin incidentes registrados en el período", 0, "0.0%"])
        curr_row = ws_resumen.max_row
        for c in range(1, 4):
            ws_resumen.cell(row=curr_row, column=c).border = border_fino
    else:
        for item in top_5:
            cat = item.get("categoria", "Incidencias Generales")
            cant = item.get("total", 0)
            pct = f"{(cant / total_t * 100):.1f}%" if total_t > 0 else "0.0%"
            ws_resumen.append([cat, cant, pct])
            curr_row = ws_resumen.max_row
            ws_resumen.cell(row=curr_row, column=1).font = font_normal
            ws_resumen.cell(row=curr_row, column=1).border = border_fino
            c2 = ws_resumen.cell(row=curr_row, column=2)
            c2.font = font_negrita
            c2.alignment = Alignment(horizontal="center")
            c2.border = border_fino
            c3 = ws_resumen.cell(row=curr_row, column=3)
            c3.font = font_normal
            c3.alignment = Alignment(horizontal="center")
            c3.border = border_fino

    ws_resumen.append([])  # Espacio

    # c) SECCIÓN: RESUMEN MENSUAL DEL PERÍODO
    row_men_sec = ws_resumen.max_row + 1
    ws_resumen.append(["RESUMEN MENSUAL DEL PERÍODO", "", "", "", "", ""])
    c_sec3 = ws_resumen.cell(row=row_men_sec, column=1)
    c_sec3.font = font_seccion_header
    c_sec3.fill = fill_seccion_header

    headers_m = ["Mes", "Año", "Total Atenciones", "Incidentes Resueltos", "% Cumplimiento SLA", "Horas Ahorradas (HU15.1)"]
    ws_resumen.append(headers_m)
    row_m_head = ws_resumen.max_row
    for c in range(1, len(headers_m) + 1):
        cell = ws_resumen.cell(row=row_m_head, column=c)
        cell.font = font_tabla_header
        cell.fill = fill_tabla_header
        cell.border = border_fino
        cell.alignment = Alignment(horizontal="center" if c > 1 else "left")

    for m in resumen_mensual:
        ws_resumen.append([
            m.get("mes", ""),
            m.get("anio", ahora.year),
            m.get("total_atenciones", 0),
            m.get("incidentes_resueltos", 0),
            f"{m.get('cumplimiento_sla_porcentaje', 100.0)}%",
            f"{m.get('horas_ahorradas', 0.0)} hrs"
        ])
        curr_row = ws_resumen.max_row
        ws_resumen.cell(row=curr_row, column=1).font = font_normal
        ws_resumen.cell(row=curr_row, column=1).border = border_fino
        for c in range(2, 7):
            cell = ws_resumen.cell(row=curr_row, column=c)
            cell.font = font_negrita if c in [3, 4, 5, 6] else font_normal
            cell.alignment = Alignment(horizontal="center")
            cell.border = border_fino

    # Ajustar ancho de columnas en Hoja 1
    for col in ws_resumen.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_resumen.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # -------------------------------------------------------------------------
    # HOJA 2: DATA DETALLADA (Detalle completo de incidencias)
    # -------------------------------------------------------------------------
    ws_data = wb.create_sheet(title="Data Detallada")

    headers_data = ["ID", "Solicitante", "Sede", "Criticidad", "Estado", "Estado SLA", "Fecha Creación"]
    ws_data.append(headers_data)

    for col_idx in range(1, len(headers_data) + 1):
        cell = ws_data.cell(row=1, column=col_idx)
        cell.font = font_tabla_header
        cell.fill = fill_tabla_header
        cell.border = border_fino
        cell.alignment = Alignment(
            horizontal="center" if headers_data[col_idx - 1] in ["ID", "Criticidad", "Estado", "Estado SLA", "Fecha Creación"] else "left",
            vertical="center"
        )

    for ticket, usuario in results:
        creador_info = estructurar_info_creador(db, ticket, usuario=usuario)
        solicitante_nombre = creador_info.nombre or ticket.correo_solicitante or (f"Usuario #{ticket.usuario_id}" if ticket.usuario_id else "Usuario")

        sede = ticket.sede or "-"
        criticidad = ticket.criticidad or "Media"
        estado = ticket.estado or "Abierto"

        sla_info = calcular_metrica_ticket(ticket, ahora)
        estado_sla = sla_info.get("estado_sla", "Normal")

        fecha_creacion_str = formatear_fecha_lima(ticket.fecha_creacion, "%Y-%m-%d %H:%M:%S")

        ws_data.append([
            ticket.id,
            solicitante_nombre,
            sede,
            criticidad,
            estado,
            estado_sla,
            fecha_creacion_str
        ])
        curr_row = ws_data.max_row
        for c in range(1, len(headers_data) + 1):
            cell = ws_data.cell(row=curr_row, column=c)
            cell.border = border_fino
            if headers_data[c - 1] in ["ID", "Criticidad", "Estado", "Estado SLA", "Fecha Creación"]:
                cell.alignment = Alignment(horizontal="center")

    # Ajustar ancho de columnas en Hoja 2
    for col in ws_data.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_data.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Asegurar que Hoja 1 ("Resumen Ejecutivo") sea la activa al abrir el archivo
    wb.active = ws_resumen

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    headers_response = {
        "Content-Disposition": 'attachment; filename="reporte_incidencias.xlsx"'
    }

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_response
    )


@dashboard_router.get("/gerencial")
def get_dashboard_gerencial(
    db: Session = Depends(get_db),
    admin_user: models.Usuario = Depends(get_current_admin_user)
):
    """
    Alias directo para el Dashboard Gerencial (HU20).
    Protegido por get_current_admin_user (403 Forbidden para usuarios no administradores).
    """
    return generar_resumen_ejecutivo(db)


@dashboard_router.get("/stats")
def get_dashboard_stats(
    db: Session = Depends(get_db),
    admin_user: models.Usuario = Depends(get_current_admin_user)
):
    """Métricas estadísticas agrupadas para el Dashboard Gerencial."""
    return generar_resumen_ejecutivo(db)
